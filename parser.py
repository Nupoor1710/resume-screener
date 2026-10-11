"""Text extraction + information extraction (skills, education, experience)."""
import io
import re
from datetime import date
from pathlib import Path

from skills import SKILLS

EDU_LEVELS = {
    4: [r"ph\.?\s?d", r"doctorate", r"doctoral"],
    3: [r"m\.?\s?tech", r"m\.?\s?sc", r"m\.?\s?s\b", r"m\.?\s?e\b", r"mba", r"mca", r"master'?s?", r"post\s?graduate"],
    2: [r"b\.?\s?tech", r"b\.?\s?e\b", r"b\.?\s?sc", r"b\.?\s?s\b", r"bca", r"b\.?\s?com", r"bachelor'?s?", r"undergraduate", r"b\.?\s?a\b"],
    1: [r"diploma", r"associate degree"],
}
EDU_NAMES = {0: "Not specified", 1: "Diploma", 2: "Bachelor's", 3: "Master's", 4: "PhD"}

MONTHS = "jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec"
RANGE_RE = re.compile(
    rf"(?:(?:{MONTHS})[a-z]*\.?\s+)?((?:19|20)\d{{2}})\s*(?:-|–|—|to)\s*"
    rf"(?:(?:{MONTHS})[a-z]*\.?\s+)?((?:19|20)\d{{2}}|present|current|now|till date)",
    re.I,
)
YEARS_RE = re.compile(r"(\d{1,2})(?:\.\d)?\s*\+?\s*(?:years?|yrs?)(?:\s+of)?(?:\s+(?:professional|relevant|work|industry))?(?:\s+experience)?", re.I)


# ---------- reading files ----------
def read_text(source, filename: str | None = None) -> str:
    """source: file path (str/Path) or bytes / file-like object (with filename)."""
    if isinstance(source, (str, Path)):
        path = Path(source)
        filename = path.name
        data = path.read_bytes()
    elif hasattr(source, "read"):
        data = source.read()
    else:
        data = source
    ext = Path(filename or "").suffix.lower()

    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" ".join(c.text for c in row.cells))
        return "\n".join(parts)
    return data.decode("utf-8", errors="ignore")


# ---------- extraction ----------
def clean(text: str) -> str:
    text = text.replace("\u2022", " ").replace("\xa0", " ")
    return re.sub(r"[ \t]+", " ", text)


def extract_skills(text: str) -> set[str]:
    low = " " + clean(text).lower() + " "
    found = set()
    for canon, aliases in SKILLS.items():
        for a in aliases:
            pat = r"(?<![a-z0-9+#])" + re.escape(a) + r"(?![a-z0-9+#])"
            if re.search(pat, low):
                found.add(canon)
                break
    return found


def extract_education_level(text: str) -> int:
    low = clean(text).lower()
    best = 0
    for level, pats in EDU_LEVELS.items():
        for p in pats:
            if re.search(r"(?<![a-z])" + p + r"(?![a-z])", low):
                best = max(best, level)
    return best


def extract_experience_years(text: str) -> float:
    """Prefer explicit 'X years' statements; otherwise sum date ranges (outside education lines)."""
    text = clean(text)
    stated = [int(m.group(1)) for m in YEARS_RE.finditer(text) if int(m.group(1)) <= 40
              and re.search(r"exp", text[m.start(): m.end() + 40], re.I)]
    if stated:
        return float(max(stated))

    edu_kw = re.compile(r"b\.?tech|m\.?tech|bachelor|master|degree|university|college|school|b\.?sc|m\.?sc|diploma|ph\.?d|cgpa|gpa", re.I)
    this_year = date.today().year
    spans = []
    for line in text.splitlines():
        if edu_kw.search(line):
            continue
        for m in RANGE_RE.finditer(line):
            s = int(m.group(1))
            e_raw = m.group(2).lower()
            e = this_year if e_raw[0].isalpha() else int(e_raw)
            if s <= e <= this_year + 1:
                spans.append((s, e))
    if not spans:
        return 0.0
    spans.sort()
    merged = [list(spans[0])]
    for s, e in spans[1:]:
        if s <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], e)
        else:
            merged.append([s, e])
    return float(sum(e - s for s, e in merged))


def extract_email(text: str) -> str:
    m = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)
    return m.group(0) if m else ""


def guess_name(text: str, fallback: str) -> str:
    for line in text.splitlines()[:8]:
        line = line.strip()
        if re.fullmatch(r"[A-Za-z][A-Za-z.'-]+(?: [A-Za-z][A-Za-z.'-]+){1,3}", line) and not re.search(
                r"resume|curriculum|vitae|engineer|developer|scientist|analyst|manager", line, re.I):
            return line.title() if line.isupper() else line
    return fallback


def parse_resume(text: str, fallback_name: str = "Unknown") -> dict:
    return {
        "name": guess_name(text, fallback_name),
        "email": extract_email(text),
        "skills": extract_skills(text),
        "education_level": extract_education_level(text),
        "experience_years": extract_experience_years(text),
        "text": text,
    }


def parse_job_description(text: str) -> dict:
    """Pull required skills, minimum experience and education out of a JD."""
    years = [int(m.group(1)) for m in re.finditer(r"(\d{1,2})\s*\+?\s*(?:years?|yrs?)", text, re.I) if int(m.group(1)) <= 25]
    return {
        "skills": extract_skills(text),
        "min_experience": float(min(years)) if years else 0.0,
        "education_level": extract_education_level(text),
        "text": text,
    }
